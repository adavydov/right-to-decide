#!/usr/bin/env python3
"""Readable Markdown to DOCX, optionally PDF. Existing same-name outputs replaced.
Subset: headings, paragraphs, lists, tables, quotes, bold, italic, links, code,
[1] references and [1]/[^1]: endnote definitions. No full CommonMark support.
"""
from pathlib import Path
import argparse
import os
import re
import subprocess
import tempfile
from urllib.parse import unquote, urlparse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
TOKEN = re.compile(r"(\[([^\]]+)\]\((?:<([^>]+)>|([^\s)]+))\)|\*\*(.+?)\*\*|(?<!\*)\*([^*]+)\*(?!\*)|\x60([^\x60]+)\x60|\[\^?(\d+)\])")
NOTE = re.compile(r"^\[\^?(\d+)\](?::)?\s+(.+)$")
LIST = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$")
FENCES = (chr(96)*3, "~~~")

def xml(tag, **attrs):
    x = OxmlElement(tag)
    for k, v in attrs.items():
        x.set(qn(k), str(v))
    return x

def configure(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Mm(22)
    s.header_distance = s.footer_distance = Mm(11)
    for name, size in {"Normal":11.5,"Title":24,"Subtitle":12,"Heading 1":19,"Heading 2":15,"Heading 3":12.5,"Quote":11.5,"List Bullet":11.5}.items():
        style = doc.styles[name]
        style.font.name, style.font.size = "Georgia", Pt(size)
        style.element.get_or_add_rPr().append(xml("w:lang", **{"w:val":"ru-RU"}))
        p = style.paragraph_format
        p.line_spacing, p.space_after, p.widow_control = 1.2, Pt(7), True
        if name.startswith("Heading") or name in ("Title","Subtitle"):
            style.font.bold = name != "Subtitle"
            style.font.color.rgb = RGBColor.from_string("243345")
            p.space_before, p.space_after = Pt(0 if name=="Title" else 16), Pt(9)
            p.keep_with_next = p.keep_together = True
    doc.styles["Quote"].paragraph_format.left_indent = Mm(8)
    doc.styles["Quote"].paragraph_format.right_indent = Mm(5)
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field, run, text = xml("w:fldSimple", **{"w:instr":"PAGE"}), xml("w:r"), xml("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    p._p.append(field)
    doc.settings.element.append(xml("w:updateFields", **{"w:val":"true"}))

def uri(target, source):
    target = target.replace("\\","/")
    if re.match(r"^[A-Za-z]:/", target):
        return Path(unquote(target)).resolve().as_uri()
    if urlparse(target).scheme or target.startswith("#"):
        return target
    file, mark, anchor = target.partition("#")
    return (source.parent/unquote(file)).resolve().as_uri() + (mark+anchor if mark else "")

def link(p, label, target, source, superscript=False):
    x = xml("w:hyperlink")
    if target.startswith("#"):
        x.set(qn("w:anchor"),target[1:])
    else:
        x.set(qn("r:id"),p.part.relate_to(uri(target,source),RT.HYPERLINK,is_external=True))
    run, props, text = xml("w:r"),xml("w:rPr"),xml("w:t")
    props.append(xml("w:color", **{"w:val":"225B83"}))
    props.append(xml("w:u", **{"w:val":"single"}))
    if superscript:
        props.append(xml("w:vertAlign", **{"w:val":"superscript"}))
    run.append(props)
    text.set(qn("xml:space"),"preserve")
    text.text = label
    run.append(text)
    x.append(run)
    p._p.append(x)

def inline(p, text, source, notes, bold=False, italic=False):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start()>pos:
            r=p.add_run(text[pos:m.start()])
            r.bold,r.italic=bold,italic
        whole,label,angled,target,strong,emphasis,code,note=m.groups()
        if label is not None:
            link(p,label,angled or target,source)
        elif strong is not None:
            inline(p,strong,source,notes,True,italic)
        elif emphasis is not None:
            inline(p,emphasis,source,notes,bold,True)
        elif code is not None:
            r=p.add_run(code)
            r.font.name,r.font.size="Consolas",Pt(10)
        elif note in notes:
            link(p,f"[{note}]",f"#note_{note}",source,True)
        else:
            r=p.add_run(whole)
            r.bold,r.italic=bold,italic
        pos=m.end()
    if pos<len(text):
        r=p.add_run(text[pos:])
        r.bold,r.italic=bold,italic

def cells(line):
    return [v.strip().replace(r"\|","|") for v in re.split(r"(?<!\\)\|",line.strip().strip("|"))]

def separator(line):
    vs=cells(line)
    return bool(vs) and all(re.fullmatch(r":?-{3,}:?",v) for v in vs)

def table(doc,rows,source,notes):
    data=[cells(row) for row in rows if not separator(row)]
    count=max(map(len,data))
    t=doc.add_table(rows=0,cols=count)
    t.style,t.autofit="Table Grid",False
    for c in t.columns:
        c.width=Mm(166/count)
    for i,values in enumerate(data):
        row=t.add_row()
        if i==0:
            row._tr.get_or_add_trPr().append(xml("w:tblHeader",**{"w:val":"true"}))
        # No fixed heights or cantSplit: long cells can continue on another page.
        for n,cell in enumerate(row.cells):
            cell.width=Mm(166/count)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.line_spacing=1.1
            p.paragraph_format.keep_with_next=False
            inline(p,values[n] if n<len(values) else "",source,notes,bold=i==0)
            for r in p.runs:
                r.font.size=Pt(10)
            if i==0:
                cell._tc.get_or_add_tcPr().append(xml("w:shd",**{"w:fill":"EEF2F5"}))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def build_docx(source,destination,title=None,subtitle=None):
    lines=source.read_text(encoding="utf-8-sig").splitlines()
    notes={m.group(1) for line in lines if (m:=NOTE.match(line.strip()))}
    doc=Document()
    configure(doc)
    doc.core_properties.title=title or next((s[2:] for s in lines if s.startswith("# ")),source.stem)
    doc.core_properties.author,doc.core_properties.subject="",subtitle or ""
    if title:
        inline(doc.add_paragraph(style="Title"),title,source,notes)
    if subtitle:
        inline(doc.add_paragraph(style="Subtitle"),subtitle,source,notes)
    i,bookmark,first=0,1,True
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            i+=1
            continue
        if line.startswith(FENCES):
            fence,block=line[:3],[]
            i+=1
            while i<len(lines) and not lines[i].strip().startswith(fence):
                block.append(lines[i])
                i+=1
            p=doc.add_paragraph()
            p.paragraph_format.line_spacing=1.0
            r=p.add_run("\n".join(block))
            r.font.name,r.font.size="Consolas",Pt(9)
            i+=1
            continue
        h=re.match(r"^(#{1,6})\s+(.+?)(?:\s+#+)?$",line)
        if h:
            level,value=min(len(h.group(1)),3),h.group(2)
            if not(first and title==value):
                style="Title" if first and level==1 and not title else f"Heading {level}"
                inline(doc.add_paragraph(style=style),value,source,notes)
            first=False
            i+=1
            continue
        if "|" in line and i+1<len(lines) and separator(lines[i+1]):
            rows=[line,lines[i+1]]
            i+=2
            while i<len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(lines[i])
                i+=1
            table(doc,rows,source,notes)
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})",line):
            doc.add_paragraph("⁂").alignment=WD_ALIGN_PARAGRAPH.CENTER
            i+=1
            continue
        note=NOTE.match(line)
        if note:
            p=doc.add_paragraph()
            p._p.append(xml("w:bookmarkStart",**{"w:id":bookmark,"w:name":f"note_{note.group(1)}"}))
            p.add_run(f"[{note.group(1)}] ")
            inline(p,note.group(2),source,notes)
            p._p.append(xml("w:bookmarkEnd",**{"w:id":bookmark}))
            bookmark+=1
            i+=1
            continue
        item=LIST.match(lines[i])
        if item:
            numbered=item.group(2)[0].isdigit()
            p=doc.add_paragraph(style="Normal" if numbered else "List Bullet")
            if numbered:
                p.paragraph_format.left_indent,p.paragraph_format.first_line_indent=Mm(7),Mm(-7)
                p.add_run(item.group(2)+"\t")
            if item.group(1):
                p.paragraph_format.left_indent=Mm(7+min(len(item.group(1)),12)*2)
            inline(p,item.group(3),source,notes)
            i+=1
            continue
        quote=line.startswith(">")
        block=[re.sub(r"^>\s?","",line) if quote else line]
        i+=1
        while i<len(lines) and lines[i].strip():
            s=lines[i].strip()
            if s.startswith(("#",)+FENCES) or LIST.match(lines[i]) or NOTE.match(s):
                break
            if s.startswith(">")!=quote or re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})",s):
                break
            if i+1<len(lines) and "|" in s and separator(lines[i+1]):
                break
            block.append(re.sub(r"^>\s?","",s) if quote else s)
            i+=1
        inline(doc.add_paragraph(style="Quote" if quote else "Normal")," ".join(block),source,notes)
    destination.parent.mkdir(parents=True,exist_ok=True)
    doc.save(destination)

def build_pdf(docx,office):
    if not office.is_file():
        raise FileNotFoundError(f"LibreOffice executable not found: {office}")
    with tempfile.TemporaryDirectory(prefix="reading-documents-lo-") as folder:
        profile=Path(folder,"profile").resolve().as_uri()
        result=subprocess.run([str(office.resolve()),f"-env:UserInstallation={profile}",
            "--headless","--convert-to","pdf","--outdir",str(docx.parent.resolve()),str(docx.resolve())],
            capture_output=True,text=True,encoding="utf-8",errors="replace",timeout=180,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name=="nt" else 0)
        if result.returncode or not docx.with_suffix(".pdf").is_file():
            raise RuntimeError(f"PDF conversion failed: {result.stdout}\n{result.stderr}")
    return docx.with_suffix(".pdf")

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown",type=Path,nargs="+")
    parser.add_argument("--output-dir",type=Path,help="Default: beside each source")
    parser.add_argument("--title",help="Title on first page; no separate cover")
    parser.add_argument("--subtitle",help="Subtitle on first page")
    parser.add_argument("--pdf",action="store_true")
    parser.add_argument("--libreoffice",type=Path,default=OFFICE)
    args=parser.parse_args()
    if (args.title or args.subtitle) and len(args.markdown)!=1:
        parser.error("--title/--subtitle require exactly one source")
    for source in args.markdown:
        source=source.resolve()
        if source.suffix.lower()!=".md" or not source.is_file():
            parser.error(f"Expected an existing Markdown file: {source}")
        dest=(args.output_dir.resolve() if args.output_dir else source.parent)/f"{source.stem}.docx"
        build_docx(source,dest,args.title,args.subtitle)
        print(dest)
        if args.pdf:
            print(build_pdf(dest,args.libreoffice))
if __name__=="__main__":
    main()

