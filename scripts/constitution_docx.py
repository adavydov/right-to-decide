"""Create/verify the complete reader copy of the canonical constitution.

Generation uses python-docx; --check uses only the standard library so CI can
verify every paragraph, table cell and hyperlink against the Markdown source.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import re
from urllib.parse import urljoin, urlsplit
from xml.etree import ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CONSTITUTION.md"
OUTPUT = ROOT / "public/manifesto/Pravo_na_reshenie_Manifest_Constitution_v1.2.1.docx"
REPOSITORY = "https://github.com/adavydov/right-to-decide/blob/main/"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def inline(value):
    pattern = re.compile(r"\[([^\]\n]+)\]\(([^\s]+)\)|\*\*(.+?)\*\*|(?<!\*)\*([^*\n]+)\*(?!\*)|`([^`\n]+)`|https?://[^\s]+")
    result = []
    cursor = 0
    for match in pattern.finditer(value):
        if match.start() > cursor:
            result.append({"text": value[cursor:match.start()]})
        if match[1] is not None:
            target = urljoin(REPOSITORY, match[2])
            parsed = urlsplit(target)
            if parsed.scheme not in {"http", "https"} or not parsed.netloc or parsed.username or parsed.password:
                raise ValueError("Unsupported constitution link: " + match[2])
            result.extend({**run, "href": target} for run in inline(match[1]))
        elif match[3] is not None:
            result.extend({**run, "bold": True} for run in inline(match[3]))
        elif match[4] is not None:
            result.extend({**run, "italic": True} for run in inline(match[4]))
        elif match[5] is not None:
            result.append({"text": match[5], "code": True})
        else:
            url = match[0].rstrip(".,;:")
            while url.endswith(")") and url.count(")") > url.count("("):
                url = url[:-1]
            result.append({"text": url, "href": url})
            if url != match[0]:
                result.append({"text": match[0][len(url):]})
        cursor = match.end()
    if cursor < len(value):
        result.append({"text": value[cursor:]})
    return result


def parse():
    raw = SOURCE.read_text("utf-8-sig")
    if not re.search(r"^Версия 1\.2\.1 ·", raw, re.MULTILINE):
        raise ValueError("Choose an explicit output version before updating the constitution")
    units = []
    paragraph = []

    def add(text, style="Normal"):
        units.append({"kind": "paragraph", "style": style, "runs": inline(text)})

    def flush():
        if paragraph:
            add(" ".join(paragraph))
            paragraph.clear()

    for line in raw.splitlines():
        if not line.strip():
            flush()
        elif line.strip() == "---":
            flush()
        elif line.startswith("|"):
            flush()
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                continue
            if not units or units[-1]["kind"] != "table":
                units.append({"kind": "table", "rows": []})
            units[-1]["rows"].append([inline(cell) for cell in cells])
        elif heading := re.match(r"^(#{1,6})\s+(.+)$", line):
            flush()
            add(heading[2], "Title" if len(heading[1]) == 1 else "Heading " + str(min(len(heading[1]) - 1, 3)))
        elif line.startswith("> "):
            flush()
            add(line[2:], "Quote")
        elif re.match(r"^\s*(?:[-*+] |\d+\. )", line):
            flush()
            add(line.strip(), "List")
        elif line.startswith(("```", "~~~", "![")):
            raise ValueError("Add an explicit renderer for this Markdown structure")
        else:
            paragraph.append(line)
    flush()
    return units


def paragraphs(units):
    for unit in units:
        if unit["kind"] == "table":
            for row in unit["rows"]:
                yield from row
        else:
            yield unit["runs"]


def verify(path=OUTPUT):
    units = parse()
    expected = ["".join(run["text"] for run in runs) for runs in paragraphs(units)]
    with ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError("Corrupt DOCX archive")
        body = ET.fromstring(archive.read("word/document.xml")).find("{" + W + "}body")
        actual = ["".join(t.text or "" for t in p.iter("{" + W + "}t")) for p in body.iter("{" + W + "}p")]
        if actual != expected:
            raise ValueError("The complete DOCX wording/order differs from CONSTITUTION.md")
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        links = {r.attrib["Target"] for r in relationships if r.attrib.get("Type") == R + "/hyperlink"}
        expected_links = {run["href"] for runs in paragraphs(units) for run in runs if "href" in run}
        if links != expected_links:
            raise ValueError("DOCX hyperlink targets differ from the canonical source")
        if len(body.findall("{" + W + "}tbl")) != sum(unit["kind"] == "table" for unit in units):
            raise ValueError("A canonical table is missing from DOCX")
    print(f"Constitution DOCX verified: {len(actual)} paragraphs/cells, {len(links)} links, complete canonical wording and order.")


def generate():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Georgia", Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Georgia", Pt(size)
        style.font.color.rgb = RGBColor.from_string("172B3A")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = "Право на решение. Манифест и конституция 1.2.1"
    doc.core_properties.subject = "Полная читательская копия CONSTITUTION.md"

    def append_runs(paragraph, runs):
        for item in runs:
            run = paragraph.add_run(item["text"])
            run.bold = item.get("bold", False)
            run.italic = item.get("italic", False)
            if item.get("code"):
                run.font.name = "Consolas"
            if item.get("href"):
                relation = paragraph.part.relate_to(item["href"], R + "/hyperlink", is_external=True)
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), relation)
                run.font.color.rgb = RGBColor.from_string("315D82")
                run.underline = True
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)

    for unit in parse():
        if unit["kind"] == "table":
            rows = unit["rows"]
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("Unequal table columns")
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, runs in zip(cells, row):
                    append_runs(cell.paragraphs[0], runs)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        if i == 0:
                            run.bold = True
                if i == 0:
                    properties = table.rows[0]._tr.get_or_add_trPr()
                    properties.append(OxmlElement("w:tblHeader"))
        else:
            style = unit["style"]
            p = doc.add_paragraph(style="Normal" if style == "List" else style)
            if style == "List":
                p.paragraph_format.left_indent = Cm(.4)
                p.paragraph_format.first_line_indent = Cm(-.4)
            append_runs(p, unit["runs"])
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    verify()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    verify() if args.check else generate()
