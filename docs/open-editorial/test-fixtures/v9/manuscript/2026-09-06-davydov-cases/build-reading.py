#!/usr/bin/env python3
"""Build only the local 7.1 reading edition; never modify site publication data."""
from __future__ import annotations
import argparse, hashlib, importlib.util, json, os, re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

HERE = Path(__file__).resolve().parent
BASE = HERE.parent / '2026-09-06-critic-revision'
CONVERTER = HERE.parent / '2026-09-05-rebuild/build_reading_documents.py'
OUTPUT = HERE / 'reading'
REPLACED = {12, 14, 17}
PARTS = [
    'Мир может развиваться без нас',
    'Как общество выращивает тех, кто способен его изменить',
    'Инженерная школа как место изобретения деятельности',
    'Кому доступно создание будущего',
    'Свобода внутри развивающей среды',
    'Образование, которого нет',
]
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI']
NOTES_HEADING = re.compile(r'(?mi)^#{1,6}\s+(?:Источники?|Примечани[ея])[^\n]*$')
DEFINITION = re.compile(r'(?m)^(?:\[\^?(\d+)\]:?|(\d+)\.)\s+(.+)$')
REFERENCE = re.compile(r'\[\^?(\d+)\](?!\()')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest().upper()
def read(path): return path.read_text(encoding='utf-8-sig')
def relative(path): return Path(os.path.relpath(path, HERE)).as_posix()

def inputs():
    return [BASE / 'prologue.md'] + [
        HERE / f'chapter-{n:02}.md' if n in REPLACED else BASE / f'chapters-v7/chapter-{n:02}.md'
        for n in range(1, 19)
    ] + [BASE / 'chapters-v7/epilogue.md']

def parse(path):
    text = read(path)
    heading = re.match(r'\A# ([^\n]+)\n', text)
    if not heading: raise ValueError(f'Missing opening title: {path.name}')
    chunks = NOTES_HEADING.split(text[heading.end():], maxsplit=1)
    body, notes = chunks[0].strip(), chunks[1].strip() if len(chunks) == 2 else ''
    ids = [m.group(1) or m.group(2) for m in DEFINITION.finditer(notes)]
    if len(ids) != len(set(ids)): raise ValueError(f'Duplicate note: {path.name}')
    refs = set(REFERENCE.findall(body))
    if refs != set(ids): raise ValueError(f'Unmatched notes in {path.name}: refs={refs}, definitions={set(ids)}')
    return {'path': path, 'title': heading.group(1).strip(), 'body': body, 'notes': notes, 'ids': ids, 'sha256': sha(path)}

def assemble():
    baseline = json.loads(read(BASE / 'reading/assembly-v7.json'))
    expected = {row['path']: row['sha256'].upper() for row in baseline['inputs']}
    sections = [parse(path) for path in inputs()]
    for section in sections:
        path = section['path']
        if path.is_relative_to(BASE):
            key = path.relative_to(BASE).as_posix()
            if section['sha256'] != expected[key]: raise ValueError(f'Preserved 7.0 source differs from manifest: {key}')
    toc = ['## Содержание', '', sections[0]['title'], '']
    for index, part in enumerate(PARTS):
        toc += [f'**Часть {ROMAN[index]}. {part}**', '']
        for section in sections[1 + index * 3:4 + index * 3]: toc += [section['title'], '']
    toc += [sections[-1]['title'], '']
    result = ['# Право на решение', '', 'Литературная редакция 7.1 · 6 сентября 2026 года', ''] + toc
    appendix, note_map = [], []
    for index, section in enumerate(sections):
        numbers = {local: str(len(note_map) + i + 1) for i, local in enumerate(section['ids'])}
        note_map += [{'input': relative(section['path']), 'local': local, 'global': int(g)} for local, g in numbers.items()]
        def renumber(match): return '[' + numbers[match.group(1)] + ']'
        body = REFERENCE.sub(renumber, section['body'])
        if 1 <= index <= 18 and (index - 1) % 3 == 0:
            p = (index - 1) // 3
            result += [f'# Часть {ROMAN[p]}. {PARTS[p]}', '']
        result += ['# ' + section['title'], '', body, '']
        if section['notes']:
            notes = DEFINITION.sub(lambda m: '[' + (m.group(1) or m.group(2)) + '] ' + m.group(3), section['notes'])
            notes = REFERENCE.sub(renumber, notes)
            appendix += ['## ' + section['title'], '', notes, '']
    result += ['# Примечания и источники', ''] + appendix
    compiled = '\n'.join(result).rstrip() + '\n'
    if re.search(r'book-memory|file://|(?<!\w)[A-Za-z]:[\\/]', compiled, re.I): raise ValueError('Private path in reader text')
    definitions = [int(m.group(1)) for m in re.finditer(r'(?m)^\[(\d+)\] ', compiled)]
    if definitions != list(range(1, len(note_map) + 1)): raise ValueError('Global note sequence is broken')
    return compiled, sections, note_map

def format_docx(path, sections):
    from docx import Document
    from docx.shared import Pt
    doc = Document(path)
    doc.core_properties.title = 'Право на решение — редакция 7.1'
    doc.core_properties.author = 'Алексей М. Давыдов'
    doc.core_properties.subject = 'Локальная читательская редакция'
    special = {sections[0]['title'], sections[-1]['title'], 'Примечания и источники'}
    in_notes = False
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Heading 1':
            if paragraph.text in special or paragraph.text.startswith('Часть '): paragraph.paragraph_format.page_break_before = True
            if paragraph.text == 'Примечания и источники': in_notes = True
        if in_notes and paragraph.style.name == 'Normal':
            for run in paragraph.runs: run.font.size = Pt(10)
    # Keep the compact budget table together, including its final total row.
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
    doc.save(path)

def verify_docx(path, sections, note_map):
    with ZipFile(path) as archive:
        tree = ET.fromstring(archive.read('word/document.xml'))
        relationships = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
    names = [node.get('{' + NS['w'] + '}name') for node in tree.findall('.//w:bookmarkStart', NS)]
    expected = [f"note_{row['global']}" for row in note_map]
    if names != expected or len(names) != len(set(names)): raise ValueError('DOCX note bookmarks differ from assembly')
    anchors = [node.get('{' + NS['w'] + '}anchor') for node in tree.findall('.//w:hyperlink', NS)]
    anchors = [v for v in anchors if v]
    if set(anchors) != set(expected): raise ValueError('DOCX internal note references incomplete')
    urls = [node.get('Target') for node in relationships if node.get('Type', '').endswith('/hyperlink')]
    source_urls = {m.group(1).strip('<>') for s in sections for m in re.finditer(r'\]\(([^)\n]+)\)', s['body'] + '\n' + s['notes'])}
    if set(urls) != source_urls: raise ValueError('DOCX external hyperlinks differ from inputs')
    visible = '\n'.join(''.join(node.text or '' for node in paragraph.findall('.//w:t', NS)) for paragraph in tree.findall('.//w:p', NS))
    raw = read(path.with_suffix('.md'))
    raw = re.sub(r'(?m)^\|[-:| ]+\|\s*$', '', raw)
    raw = re.sub(r'(?m)^(?:-{3,}|\*{3,}|_{3,})\s*$', '⁂', raw)
    raw = re.sub(r'(?m)^#{1,6}\s+', '', raw)
    raw = re.sub(r'\[([^\]]+)\]\([^)\n]+\)', r'\1', raw)
    raw = re.sub(r'(?m)^\|(.+)\|\s*$', lambda m: m.group(1).replace('|', ' '), raw)
    raw = raw.replace('*', '').replace(chr(96), '')
    norm = lambda value: re.sub(r'\s+', ' ', value).strip()
    if norm(raw) != norm(visible):
        import difflib
        delta = '\n'.join(list(difflib.unified_diff(norm(raw).split(), norm(visible).split(), n=5))[:35])
        raise ValueError('DOCX visible text differs from compiled Markdown:\n' + delta)
    for s in sections:
        if s['title'] not in visible: raise ValueError('Missing DOCX title: ' + s['title'])
    return {'completeVisibleTextMatchesMarkdown': True, 'noteBookmarks': len(names), 'internalNoteLinks': len(anchors), 'externalHyperlinks': len(urls), 'inputTitlesPresent': len(sections)}


def verify_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    clean = []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ''
        if not text.strip(): raise ValueError(f'Empty PDF page: {number}')
        # LibreOffice emits the footer before the body in this PDF text layer.
        # Permit either edge, but remove only this page's exact printed number.
        text = re.sub(r'^\s*' + str(number) + r'\s*\n', '', text, count=1)
        text = re.sub(r'\n\s*' + str(number) + r'\s*$', '', text, count=1)
        clean.append(text.replace('\u00ad', ''))
    with ZipFile(path.with_suffix('.docx')) as archive:
        tree = ET.fromstring(archive.read('word/document.xml'))
    word = '\n'.join(''.join(node.text or '' for node in paragraph.findall('.//w:t', NS)) for paragraph in tree.findall('.//w:p', NS))
    norm = lambda value: re.sub(r'\s+', '', value)
    if norm(word) != norm('\n'.join(clean)): raise ValueError('PDF text differs from DOCX after whitespace and exact page-footer removal')
    return {'pdfPages': len(reader.pages), 'pdfNonemptyPages': len(reader.pages), 'completePdfTextMatchesDocx': True}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf', action='store_true', help='Also convert DOCX through installed LibreOffice')
    args = parser.parse_args()
    compiled, sections, note_map = assemble()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    markdown = OUTPUT / 'right-to-decide-v7.1.md'
    markdown.write_text(compiled, encoding='utf-8')
    spec = importlib.util.spec_from_file_location('reading_documents', CONVERTER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    # The legacy converter displays Markdown emphasis inside hyperlink labels literally.
    # Clean those labels in this local adapter; preserve URLs and the original converter.
    original_link = module.link
    def reader_link(paragraph, label, target, source, superscript=False):
        label = re.sub(r'\*+([^*]+)\*+', r'\1', label)
        original_link(paragraph, label, target, source, superscript)
    module.link = reader_link
    docx = markdown.with_suffix('.docx')
    module.build_docx(markdown, docx, None, None)
    format_docx(docx, sections)
    checks = verify_docx(docx, sections, note_map)
    outputs = [markdown, docx]
    if args.pdf:
        pdf = module.build_pdf(docx, module.OFFICE)
        checks.update(verify_pdf(pdf))
        outputs.append(pdf)
    for s in sections:
        if sha(s['path']) != s['sha256']: raise RuntimeError('Source changed during assembly')
    manifest = {
        'status': 'local-reading-edition', 'edition': '7.1', 'created': '2026-09-06', 'published': False,
        'literaryAcceptance': 'See separate editorial review; this file validates assembly only',
        'baselineManifest': relative(BASE / 'reading/assembly-v7.json'), 'baselineManifestSha256': sha(BASE / 'reading/assembly-v7.json'),
        'inputs': [{'order': i, 'path': relative(s['path']), 'title': s['title'], 'sha256': s['sha256'],
                    'edition': '7.1' if s['path'].parent == HERE else '7.0'} for i, s in enumerate(sections)],
        'parts': PARTS, 'noteMap': note_map, 'noteCount': len(note_map), 'markdownSha256': sha(markdown),
        'builder': 'build-reading.py', 'builderSha256': sha(Path(__file__)), 'converter': relative(CONVERTER), 'converterSha256': sha(CONVERTER),
        'outputs': [{'name': p.name, 'bytes': p.stat().st_size, 'sha256': sha(p)} for p in outputs],
        'checks': {'preservedV7Inputs': 17, 'allInputHashesStable': True, 'globalNoteDefinitionsSequential': True, 'privatePathsAbsent': True, **checks},
        'command': 'python build-reading.py' + (' --pdf' if args.pdf else ''),
    }
    (OUTPUT / 'assembly-v7.1.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'edition': '7.1', 'inputs': len(sections), 'notes': len(note_map), 'outputs': [p.name for p in outputs], 'checks': checks}, ensure_ascii=False))

if __name__ == '__main__': main()
