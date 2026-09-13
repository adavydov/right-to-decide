"""Build local reading artifacts for edition 9.0; never accept or publish a manuscript."""
from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import sys
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent / "technical"))
import manuscript_v9
import release_metadata
from datetime import date

ROOT = Path(__file__).resolve().parents[2]
BASE = Path("manuscript/2026-09-06-corpus-rebuild")
OUTPUT = BASE / "reading"
ADAPTER = Path("manuscript/2026-09-06-davydov-cases/build-reading.py")
CONVERTER = Path("manuscript/2026-09-05-rebuild/build_reading_documents.py")

def load(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def inputs():
    return [BASE / "prologue.md"] + [BASE / f"chapters/chapter-{n:02d}.md" for n in range(1, 19)] + [BASE / "epilogue.md"]

def supplements():
    selection = json.loads((ROOT / BASE / "supplements.json").read_text("utf-8-sig"))
    entries = selection.get("appendices")
    if not isinstance(entries, list):
        raise ValueError("supplements.json must explicitly select appendices")
    result, seen = [], set()
    for entry in entries:
        relative = Path(entry["path"])
        path = (ROOT / BASE / relative).resolve()
        ident = entry["id"]
        if (relative.is_absolute() or not path.is_relative_to((ROOT / BASE).resolve())
                or path.suffix != ".md" or not re.fullmatch(r"appendix-[a-z0-9-]+", ident)
                or ident in seen):
            raise ValueError("Unsafe or duplicate appendix selection")
        seen.add(ident)
        result.append({"path": relative.as_posix(), "id": ident})
    return result

def assemble(metadata):
    legacy = load("reading_v71_adapter", ADAPTER)
    sections, note_map = [], []
    selected_supplements = supplements()
    for order, relative in enumerate(inputs() + [BASE / item["path"] for item in selected_supplements]):
        path = ROOT / relative
        ident = ("prologue" if order == 0 else "epilogue" if order == 19 else
                 f"chapter-{order:02d}" if order < 19 else selected_supplements[order - 20]["id"])
        manuscript_v9.parse(path.read_bytes(), ident)
        section = legacy.parse(path)
        section["id"] = ident
        section["kind"] = "appendix" if order >= 20 else "literary"
        sections.append(section)
    toc = ["## Содержание", "", sections[0]["title"], ""]
    for index, part in enumerate(legacy.PARTS):
        toc += [f"**Часть {legacy.ROMAN[index]}. {part}**", ""]
        for section in sections[1 + index * 3:4 + index * 3]:
            toc += [section["title"], ""]
    toc += [sections[19]["title"], ""]
    for section in sections[20:]:
        toc += [section["title"], ""]
    result = ["# Право на решение", "", "Литературная редакция 9.0 · " + release_metadata.display_date(metadata["releaseDate"]), ""] + toc
    appendix = []
    for index, section in enumerate(sections):
        relative = section["path"].relative_to(ROOT / BASE).as_posix()
        numbers = {local: str(len(note_map) + i + 1) for i, local in enumerate(section["ids"])}
        note_map += [{"input": relative, "local": local, "global": int(global_)} for local, global_ in numbers.items()]
        def renumber(match):
            return "[" + numbers[match.group(1)] + "]"
        body = legacy.REFERENCE.sub(renumber, section["body"])
        if 1 <= index <= 18 and (index - 1) % 3 == 0:
            part = (index - 1) // 3
            result += [f"# Часть {legacy.ROMAN[part]}. {legacy.PARTS[part]}", ""]
        result += ["# " + section["title"], "", body, ""]
        if section["notes"]:
            notes = legacy.DEFINITION.sub(lambda m: "[" + (m.group(1) or m.group(2)) + "] " + m.group(3), section["notes"])
            notes = legacy.REFERENCE.sub(renumber, notes)
            appendix += ["## " + section["title"], "", notes, ""]
    result += ["# Примечания и источники", ""] + appendix
    compiled = "\n".join(result).rstrip() + "\n"
    if re.search(r"book-memory|file://|(?<!\w)[A-Za-z]:[\\/]", compiled, re.I):
        raise ValueError("Private path in reading text")
    definitions = [int(m.group(1)) for m in re.finditer(r"(?m)^\[(\d+)\] ", compiled)]
    if definitions != list(range(1, len(note_map) + 1)):
        raise ValueError("Global note sequence is broken")
    return compiled, sections, note_map, legacy

def verify_docx_v9(path, sections, note_map):
    """Compare every word/link while accounting for genuine Word bullet styles."""
    from docx import Document
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
        relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    names = [node.get("{" + ns["w"] + "}name") for node in tree.findall(".//w:bookmarkStart", ns)]
    expected = [f"note_{row['global']}" for row in note_map]
    if names != expected or len(names) != len(set(names)):
        raise ValueError("DOCX note bookmarks differ from assembly")
    anchors = [node.get("{" + ns["w"] + "}anchor") for node in tree.findall(".//w:hyperlink", ns)]
    anchors = [value for value in anchors if value]
    if set(anchors) != set(expected):
        raise ValueError("DOCX internal note references incomplete")
    urls = [node.get("Target") for node in relationships if node.get("Type", "").endswith("/hyperlink")]
    source_urls = {m.group(1).strip("<>") for section in sections
                   for m in re.finditer(r"\]\(([^)\n]+)\)", section["body"] + "\n" + section["notes"])}
    if set(urls) != source_urls:
        raise ValueError("DOCX external hyperlinks differ from inputs")
    visible = "\n".join("".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
                        for paragraph in tree.findall(".//w:p", ns))
    raw = path.with_suffix(".md").read_text(encoding="utf-8-sig")
    raw = re.sub(r"(?m)^\|[-:| ]+\|\s*$", "", raw)
    raw = re.sub(r"(?m)^(?:-{3,}|\*{3,}|_{3,})\s*$", "⁂", raw)
    raw = re.sub(r"(?m)^#{1,6}\s+", "", raw)
    # The converter stores unordered markers as List Bullet style, not w:t text.
    # Strip only Markdown list syntax, and require the same Word list count.
    raw, bullets = re.subn(r"(?m)^([ \t]*)[-+*][ \t]+", r"\1", raw)
    docx_bullets = sum(p.style.name == "List Bullet" for p in Document(path).paragraphs)
    if bullets != docx_bullets:
        raise ValueError("DOCX bullet list count differs from Markdown")
    raw = re.sub(r"\[([^\]]+)\]\([^)\n]+\)", r"\1", raw)
    raw = re.sub(r"(?m)^\|(.+)\|\s*$", lambda m: m.group(1).replace("|", " "), raw)
    raw = raw.replace("*", "").replace(chr(96), "")
    norm = lambda value: re.sub(r"\s+", " ", value).strip()
    if norm(raw) != norm(visible):
        import difflib
        delta = "\n".join(list(difflib.unified_diff(norm(raw).split(), norm(visible).split(), n=5))[:35])
        raise ValueError("DOCX visible text differs from compiled Markdown:\n" + delta)
    for section in sections:
        if section["title"] not in visible:
            raise ValueError("Missing DOCX title: " + section["title"])
    return {"completeVisibleTextMatchesMarkdown": True, "noteBookmarks": len(names),
            "internalNoteLinks": len(anchors), "externalHyperlinks": len(urls),
            "inputTitlesPresent": len(sections), "docxBulletLists": docx_bullets}


def verify_pdf_v9(path, expected_bullets):
    """Retain full text equality; remove only exact page numbers and list glyphs."""
    from pypdf import PdfReader
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    reader, clean, bullet_count = PdfReader(path), [], 0
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if not text.strip():
            raise ValueError(f"Empty PDF page: {number}")
        text = re.sub(r"^\s*" + str(number) + r"\s*\n", "", text, count=1)
        text = re.sub(r"\n\s*" + str(number) + r"\s*$", "", text, count=1)
        # Word Symbol-font bullets may be extracted by LibreOffice as U+F0B7.
        text, count = re.subn(r"(?m)^[ \t]*[\u2022\uf0b7][ \t]*", "", text)
        bullet_count += count
        clean.append(text.replace("\u00ad", ""))
    if bullet_count != expected_bullets:
        raise ValueError(f"PDF bullet list count differs from DOCX: {bullet_count} != {expected_bullets}")
    with ZipFile(path.with_suffix(".docx")) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    word = "\n".join("".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
                     for paragraph in tree.findall(".//w:p", ns))
    norm = lambda value: re.sub(r"\s+", "", value)
    if norm(word) != norm("\n".join(clean)):
        raise ValueError("PDF text differs from DOCX after whitespace, exact footer and list-marker removal")
    return {"pdfPages": len(reader.pages), "pdfNonemptyPages": len(reader.pages),
            "completePdfTextMatchesDocx": True, "pdfBulletLists": bullet_count}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    metadata, metadata_record = release_metadata.read_metadata(ROOT, BASE / "release-metadata.json")
    compiled, sections, note_map, legacy = assemble(metadata)
    (ROOT / OUTPUT).mkdir(parents=True, exist_ok=True)
    markdown = ROOT / OUTPUT / "right-to-decide-v9.0.md"
    markdown.write_text(compiled, encoding="utf-8")
    converter = load("reading_v9_converter", CONVERTER)
    original_link = converter.link
    def reader_link(paragraph, label, target, source, superscript=False):
        label = re.sub(r"\*+([^*]+)\*+", r"\1", label)
        original_link(paragraph, label, target, source, superscript)
    converter.link = reader_link
    docx = markdown.with_suffix(".docx")
    converter.build_docx(markdown, docx, None, None)
    legacy.format_docx(docx, sections[:20])
    from docx import Document
    doc = Document(docx)
    from docx.oxml.ns import qn
    table_leads = 0
    for paragraph in doc.paragraphs:
        following = paragraph._p.getnext()
        if paragraph.text in {item["title"] for item in sections[20:]} and paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
        if following is not None and following.tag == qn("w:tbl"):
            paragraph.paragraph_format.keep_with_next = True
            table_leads += 1
    doc.core_properties.title = "Право на решение — редакция 9.0"
    doc.core_properties.subject = "Читательская сборка; литературная приёмка оформляется отдельно"
    doc.save(docx)
    checks = verify_docx_v9(docx, sections, note_map)
    checks["tableLeadParagraphsKeptWithNext"] = table_leads
    outputs = [markdown, docx]
    if args.pdf:
        pdf = converter.build_pdf(docx, converter.OFFICE)
        checks.update(verify_pdf_v9(pdf, checks["docxBulletLists"]))
        outputs.append(pdf)
    for section in sections:
        if sha(section["path"]) != section["sha256"].lower():
            raise ValueError("Source changed during reading assembly")
    manifest = {
        "schemaVersion": 1, "status": "local-reading-edition", "edition": "9.0",
        "created": date.today().isoformat(), "releaseDate": metadata["releaseDate"], "published": False,
        "releaseMetadata": metadata_record, "releaseMetadataReader": {"path": (BASE / "technical/release_metadata.py").as_posix(), "sha256": sha(ROOT / BASE / "technical/release_metadata.py")},
        "literaryAcceptance": "Not established by this technical assembly",
        "inputs": [{"order": i, "id": s["id"], "kind": s["kind"], "path": s["path"].relative_to(ROOT / BASE).as_posix(),
                    "title": s["title"], "sha256": s["sha256"].lower(), "edition": "9.0"} for i, s in enumerate(sections)],
        "parts": legacy.PARTS, "noteMap": note_map, "noteCount": len(note_map),
        "literaryNoteCount": sum(row["input"] not in {s["path"] for s in supplements()} for row in note_map),
        "supplementsSelection": {"path": (BASE / "supplements.json").as_posix(), "sha256": sha(ROOT / BASE / "supplements.json")},
        "parser": "manuscript/2026-09-06-corpus-rebuild/technical/manuscript_v9.py",
        "parserSha256": sha(ROOT / BASE / "technical/manuscript_v9.py"),
        "markdownSha256": sha(markdown),
        "builder": "manuscript/2026-09-06-corpus-rebuild/build-reading.py", "builderSha256": sha(Path(__file__)),
        "converter": CONVERTER.as_posix(), "converterSha256": sha(ROOT / CONVERTER),
        "adapter": ADAPTER.as_posix(), "adapterSha256": sha(ROOT / ADAPTER),
        "outputs": [{"name": p.name, "bytes": p.stat().st_size, "sha256": sha(p)} for p in outputs],
        "checks": {"allInputHashesStable": True, "globalNoteDefinitionsSequential": True,
                   "privatePathsAbsent": True, **checks},
        "command": "python manuscript/2026-09-06-corpus-rebuild/build-reading.py" + (" --pdf" if args.pdf else ""),
    }
    (ROOT / OUTPUT / "assembly-v9.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"edition": "9.0", "inputs": len(sections), "literaryInputs": 20, "appendices": len(sections) - 20, "notes": len(note_map),
                      "published": False, "checks": checks}, ensure_ascii=False))

if __name__ == "__main__":
    main()
