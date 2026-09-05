from pathlib import Path
import re, sys, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
sys.stdout.reconfigure(encoding='utf-8')
base=Path(r'C:\Users\Алексей\Documents\ChatGPT\Ньютон книги\docs\editorial\agent-system-2026-09-05')
src=base/'editorial-agent-architecture.md'
doc=Document()
sec=doc.sections[0]
sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68)
sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
normal=doc.styles['Normal']
normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',24),('Heading 1',17),('Heading 2',13)]:
    s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('123846')
    s.paragraph_format.space_before=Pt(15);s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Право на решение · Проект агентной редакции · ')
r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('62717A')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def add_inline(p,text):
    text=text.replace(chr(96), "")
    pat=r'(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*)'
    pos=0
    for m in re.finditer(pat,text):
        p.add_run(text[pos:m.start()])
        if m.group(2):
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m.group(3),RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr')
            col=OxmlElement('w:color');col.set(qn('w:val'),'17607D');props.append(col)
            u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u)
            run.append(props);t=OxmlElement('w:t');t.text=m.group(2);run.append(t);h.append(run);p._p.append(h)
        else:p.add_run(m.group(4)).bold=True
        pos=m.end()
    p.add_run(text[pos:])
lines=src.read_text(encoding='utf-8-sig').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            ln=lines[i].strip()
            if not re.match(r'^\|[\s:|\-]+\|$',ln):rows.append([c.strip() for c in ln.strip('|').split('|')])
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for j,txt in enumerate(rows[0]):add_inline(table.rows[0].cells[j].paragraphs[0],txt)
        for vals in rows[1:]:
            cells=table.add_row().cells
            for j,txt in enumerate(vals):add_inline(cells[j].paragraphs[0],txt)
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in table.rows:
            keep=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(keep)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for run in p.runs:run.font.size=Pt(9)
        doc.add_paragraph()
        continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title');add_inline(p,line[2:])
    elif line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1');add_inline(p,line[3:])
    elif line.startswith('### '):
        p=doc.add_paragraph(style='Heading 2');add_inline(p,line[4:])
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');add_inline(p,line[2:])
    elif line.startswith('> '):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.22)
        add_inline(p,line[2:])
        for r in p.runs:r.italic=True
    else:
        p=doc.add_paragraph();add_inline(p,line)
    i+=1
doc.core_properties.title='Право на решение: проект агентной редакции'
doc.core_properties.subject='Роли, общая память, управление, проверки и пилот'
doc.core_properties.author='Редакционная записка для авторов монографии'
target=base/'editorial-agent-architecture.docx'
doc.save(target)
check=Document(target)
links=sum(1 for rel in check.part.rels.values() if rel.reltype==RT.HYPERLINK)
text=src.read_text(encoding='utf-8-sig')
print(json.dumps({'docx':str(target),'bytes':target.stat().st_size,'paragraphs':len(check.paragraphs),'tables':len(check.tables),'external_links':links,'markdown_words':len(text.split()),'sections':sum(l.startswith('## ') for l in lines)},ensure_ascii=False))

